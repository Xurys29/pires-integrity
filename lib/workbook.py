"""Read uploaded XLSX without evaluating formulas or following external links."""
import sys, json, datetime, zipfile
import openpyxl
p=sys.argv[1]
with zipfile.ZipFile(p) as z:
    if len(z.infolist())>2000 or sum(x.file_size for x in z.infolist())>100_000_000:
        raise ValueError('Planilha excede os limites de descompactação')
w=openpyxl.load_workbook(p,read_only=True,data_only=False,keep_links=False)
def scalar(v):
    if isinstance(v,(datetime.datetime,datetime.date)): return v.isoformat()
    return v if v is None or isinstance(v,(str,int,float,bool)) else str(v)
result=[]
for s in w:
    if s.max_row>50000 or s.max_column>300: raise ValueError('Aba excede 50 mil linhas ou 300 colunas')
    rows=[]
    for row in s.iter_rows():
        values=[scalar(c.value) for c in row]
        while values and values[-1] is None: values.pop()
        if values: rows.append({'row':row[0].row,'values':values})
    result.append({'name':s.title,'rows':rows,'maxRow':s.max_row,'maxColumn':s.max_column})
json.dump({'sheets':result},sys.stdout,ensure_ascii=False,allow_nan=False)
